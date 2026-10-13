"""Build the current AI6130 assignment report draft as a DOCX file."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent.parent
OUTPUT = ROOT / "AI6130_Assignment1_Report_Draft.docx"
NAVY = "17324D"
TEAL = "246A73"
PALE = "EAF1F4"
WHITE = "FFFFFF"

BASELINE = {
    "TinyLlama-1.1B-v1.1": [
        ("2.11", "9.4", "22.39"),
        ("2.85", "11.6", "24.56"),
        ("16.52", "34.5", "47.92"),
    ],
    "Qwen2.5-3B-Instruct": [
        ("63.10", "88.5", "71.34"),
        ("57.43", "85.3", "67.30"),
        ("63.94", "82.9", "77.12"),
    ],
    "DeepSeek-R1-Distill-Qwen-1.5B": [
        ("45.21", "94.5", "47.84"),
        ("49.69", "86.8", "57.28"),
        ("30.75", "50.0", "61.49"),
    ],
    "Qwen3-4B-Instruct-2507": [
        ("41.97", "60.4", "69.46"),
        ("53.56", "70.7", "75.79"),
        ("68.53", "95.4", "71.84"),
    ],
}


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str | None = None,
                  size: float = 8.2) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    for index, line in enumerate(text.split("\n")):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        run.bold = bold
        run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def add_table(document: Document, headers: list[str], rows: list[list[str]],
              widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, label in enumerate(headers):
        shade(table.rows[0].cells[index], NAVY)
        set_cell_text(table.rows[0].cells[index], label, bold=True, color=WHITE, size=8)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, value in enumerate(values):
            if row_index % 2 == 1:
                shade(cells[column], PALE)
            set_cell_text(cells[column], value, size=7.7 if len(headers) > 3 else 8.1)
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run(text)


def add_body(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.add_run(text)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.18)
    paragraph.paragraph_format.first_line_indent = Inches(-0.12)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def setup_document() -> Document:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.58)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)
    section.header_distance = Inches(0.25)
    section.footer_distance = Inches(0.25)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.03
    for name, size, color in (("Title", 25, NAVY), ("Heading 1", 15, NAVY),
                              ("Heading 2", 10.5, TEAL)):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(7 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(4)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header.add_run("AI6130  |  ASSIGNMENT 1")
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(TEAL)
    add_page_number(section.footer.paragraphs[0])
    return document


def add_cover(document: Document) -> None:
    for _ in range(4):
        document.add_paragraph()
    eyebrow = document.add_paragraph()
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = eyebrow.add_run("AI6130  /  LARGE LANGUAGE MODELS")
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Benchmark Evaluation\nwith Open LLMs")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(8)
    run = subtitle.add_run("Baseline Evaluation and Prompt Engineering")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(TEAL)

    document.add_paragraph()
    info = document.add_table(rows=3, cols=2)
    info.alignment = WD_TABLE_ALIGNMENT.CENTER
    info.style = "Table Grid"
    for row, (label, value) in zip(info.rows, [
        ("Student", "[Your name]"),
        ("Student ID", "[Your student ID]"),
        ("Submission date", "[Date]"),
    ]):
        shade(row.cells[0], PALE)
        set_cell_text(row.cells[0], label, bold=True, color=NAVY, size=10)
        set_cell_text(row.cells[1], value, size=10)
        row.cells[0].width = Inches(1.45)
        row.cells[1].width = Inches(3.3)

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(20)
    run = note.add_run("Draft for review  |  Improved evaluation is still running")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(TEAL)
    document.add_page_break()


def add_methods_page(document: Document) -> None:
    add_heading(document, "1. Objective and Experimental Setup")
    add_body(
        document,
        "This study evaluates four open language models on three multiple-choice "
        "tasks from Open-Style/Open-LLM-Benchmark, then tests whether a constrained "
        "answer-format prompt improves evaluation outcomes. The benchmark provides "
        "an answerKey label for each item; model generations are mapped to labels "
        "before accuracy is calculated.",
    )
    add_table(document, ["Task", "Items", "Answer labels"], [
        ["CommonsenseQA", "710", "A-E"],
        ["OpenbookQA", "491", "A-D"],
        ["PiQA", "696", "A-B"],
    ], [2.0, 1.0, 1.4])

    add_heading(document, "Models and Inference", 2)
    add_table(document, ["Model", "Input format", "Max new tokens"], [
        ["TinyLlama/TinyLlama_v1.1", "Plain text", "64"],
        ["Qwen/Qwen2.5-3B-Instruct", "Chat template", "64"],
        ["DeepSeek-R1-Distill-Qwen-1.5B", "Chat template", "1,024"],
        ["Qwen/Qwen3-4B-Instruct-2507", "Chat template", "192"],
    ], [3.0, 1.55, 1.1])
    add_body(
        document,
        "All runs use greedy decoding (do_sample=False). The generation budget is "
        "held constant between baseline and improved runs for each model, while it "
        "differs across models to accommodate their output behavior. Instruction-tuned "
        "models use their tokenizer-provided chat template; TinyLlama is prompted as "
        "a plain-text completion model. No few-shot examples are used.",
    )

    add_heading(document, "Prompt Design", 2)
    add_body(document, "Baseline follows the supplied example: question, labeled options, and an open-ended answer cue.")
    code = document.add_paragraph()
    code.paragraph_format.left_indent = Inches(0.18)
    code.paragraph_format.space_after = Pt(4)
    run = code.add_run("{question}\nA: {option A}\n...\nAnswer: ")
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)
    add_body(
        document,
        "The improved prompt preserves the question and option text, then adds: "
        "“Choose the best answer from the options above. Output exactly one label "
        "from [valid labels] and nothing else.” The valid labels are generated from "
        "the current item (A-E, A-D, or A-B), followed by the same Answer cue. This "
        "targets answer readability and parser compatibility, not additional task knowledge.",
    )
    document.add_page_break()


def baseline_rows() -> list[list[str]]:
    rows = []
    for model, task_values in BASELINE.items():
        formatted = []
        for accuracy, parse_rate, parsed_accuracy in task_values:
            formatted.append(
                f"Acc. {accuracy}%\nParse {parse_rate}%\nParsed acc. {parsed_accuracy}%"
            )
        rows.append([model, *formatted])
    return rows


def add_results_page(document: Document) -> None:
    add_heading(document, "2. Results")
    add_body(
        document,
        "Table 2 reports the completed baseline results. Overall accuracy counts "
        "unparsed generations as incorrect. Parse rate is the proportion mapped to "
        "a valid answer label; parsed accuracy is calculated only over those mapped outputs.",
    )
    add_table(document, ["Model", "CommonsenseQA", "OpenbookQA", "PiQA"], baseline_rows(),
              [2.05, 1.55, 1.55, 1.25])
    caption = document.add_paragraph("Table 2. Baseline accuracy and output parseability (%).")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(5)
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8)

    add_heading(document, "Improved Prompt Results", 2)
    add_body(
        document,
        "The full improved evaluation is in progress. Populate Table 3 and the "
        "comparison discussion after all 12 model-task runs complete and the final "
        "score file is available at results/improved/scores/improved_final.csv. "
        "Do not treat partial JSONL files as final results.",
    )
    placeholder_rows = []
    for model in BASELINE:
        placeholder_rows.append([model, "Pending", "Pending", "Pending"])
    add_table(document, ["Model", "CommonsenseQA", "OpenbookQA", "PiQA"], placeholder_rows,
              [2.05, 1.55, 1.55, 1.25])
    caption = document.add_paragraph("Table 3. Improved prompt results (to be completed).")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.italic = True
        run.font.size = Pt(8)

    add_heading(document, "Baseline Observations", 2)
    add_bullet(document, "Qwen2.5 has the most consistent baseline parse rates (82.9-88.5%) across the three tasks.")
    add_bullet(document, "Qwen3 reaches 68.53% overall accuracy on PiQA, but its CommonsenseQA parse rate is 60.4%; overall accuracy there should be interpreted with this limitation in view.")
    add_bullet(document, "TinyLlama has low parse rates (9.4-34.5%). Its outputs often continue the prompt or do not provide an extractable choice, so overall accuracy understates performance on the small parsed subset and is not a clean ability estimate.")
    add_bullet(document, "DeepSeek's PiQA parse rate is 50.0%, compared with 86.8-94.5% on the other tasks; its total accuracy is therefore especially sensitive to output format and answer completion.")
    document.add_page_break()


def add_discussion_page(document: Document) -> None:
    add_heading(document, "3. Discussion and Limitations")
    add_body(
        document,
        "The baseline results show that model size or reasoning-oriented training alone "
        "does not guarantee strong accuracy on these tasks under a short, open-ended "
        "answer prompt. Output behavior matters: an answer may be correct in the "
        "generation but not recoverable as a unique label. The improved prompt is "
        "designed to reduce that format failure. Its effectiveness must be judged using "
        "both overall accuracy and parse rate, alongside parsed accuracy, once the full "
        "matrix is complete.",
    )
    add_heading(document, "Interpretation After Improved Runs", 2)
    add_bullet(document, "If parse rate rises while parsed accuracy is stable, the main gain is more reliable answer formatting rather than a clear increase in task-solving accuracy.")
    add_bullet(document, "If both overall and parsed accuracy rise, the prompt may also have helped models focus on selecting among the provided options.")
    add_bullet(document, "If a model remains strongly concentrated on one label, report its prediction distribution and compare it with the gold-label distribution; accuracy alone can hide a degenerate predictor.")

    add_heading(document, "Limitations", 2)
    add_bullet(document, "The evaluation uses one dataset split and one greedy generation per question; results may vary with sampling or alternative prompt wording.")
    add_bullet(document, "The parser is deliberately conservative when outputs mention multiple choices. This avoids crediting a distractor mention as the selected answer, but may classify some answer-bearing explanations as unparsed.")
    add_bullet(document, "Model-card results use different datasets, prompts, decoding settings, and metrics. They provide capability context, not directly comparable target scores for this experiment.")
    add_bullet(document, "The current comparison changes the prompt, while task samples and per-model generation settings remain fixed. Any conclusions should be limited to this setup.")

    add_heading(document, "References", 2)
    references = [
        "Open-Style, Open-LLM-Benchmark dataset: https://huggingface.co/datasets/Open-Style/Open-LLM-Benchmark",
        "TinyLlama model card: https://huggingface.co/TinyLlama/TinyLlama_v1.1",
        "Qwen2.5-3B-Instruct model card: https://huggingface.co/Qwen/Qwen2.5-3B-Instruct",
        "DeepSeek-R1-Distill-Qwen-1.5B model card: https://huggingface.co/deepseek-ai/DeepSeek-R1-Distill-Qwen-1.5B",
        "Qwen3-4B-Instruct-2507 model card: https://huggingface.co/Qwen/Qwen3-4B-Instruct-2507",
    ]
    for reference in references:
        add_bullet(document, reference)

    note = document.add_paragraph()
    note.paragraph_format.space_before = Pt(6)
    run = note.add_run("Before submission: fill in student details, replace the improved-results placeholders, update the comparison and conclusion from the final scores, and verify the exported page count (maximum five pages including this cover).")
    run.bold = True
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(TEAL)


def main() -> None:
    document = setup_document()
    add_cover(document)
    add_methods_page(document)
    add_results_page(document)
    add_discussion_page(document)
    document.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    main()